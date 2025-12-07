from typing import List

import openpyxl
from openpyxl import utils
from openpyxl.styles import PatternFill

from batch_data_types.xslx_types import XlsxFormatOption


class XlsxUtil:

    @staticmethod
    def read_xlsx(file_path):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        data = []
        for row in sheet.iter_rows(values_only=True):
            data.append(list(row))
        return data

    @staticmethod
    def write_xlsx(file_path,
                   data,
                   xlsx_format_options: List[XlsxFormatOption] = None
                   ):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        if xlsx_format_options:
            for o in xlsx_format_options:
                format_key = o.format_key
                format_key_name = o.format_key_name
                start_index = o.start_index or 2
                is_success_fn = o.is_success_fn

                # 如果没有设置format_key，那么从format_key_name中寻找
                if not format_key and format_key_name:
                    # 根据列名找到对应的列字母
                    for col in range(1, sheet.max_column + 1):
                        cell_value = sheet.cell(row=1, column=col).value
                        if cell_value == format_key_name:
                            format_key = utils.get_column_letter(col)
                            break

                # 如果没有找到格式化列，跳过
                if not format_key:
                    continue

                for row in range(start_index, sheet.max_row + 1):
                    cell = f"{format_key}{row}"
                    cell_value = sheet[cell].value
                    if is_success_fn:
                        if is_success_fn(cell_value):
                            sheet[cell].fill = PatternFill(fill_type='solid', fgColor='90EE90')  # 背景色为浅绿色
                        else:
                            sheet[cell].fill = PatternFill(fill_type='solid', fgColor='FFB6C1')  # 背景色为浅红色
        workbook.save(file_path)
